"""Stage 1: Load raw data from various sources into a list of Documents
(for unstructured) or a pandas DataFrame (for tabular).

Sources:
  * file: csv/tsv/xlsx/xls/pdf/docx/txt
  * paste: a single string
  * huggingface: by dataset id (and optional split/config)

Each loader returns either:
  * list[Document] — for unstructured / per-file output
  * pandas DataFrame — for tabular sources, with column names preserved
                       so the user can pick which columns matter

The structure.py stage handles tabular → list[Document].
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Union

import pandas as pd

from .segmenters import Document

# Result type from a loader: either documents or a tabular DataFrame
LoadResult = Union[list[Document], pd.DataFrame]


SUPPORTED_FILE_EXTS = {".csv", ".tsv", ".xlsx", ".xls", ".pdf", ".docx", ".txt", ".json", ".jsonl"}


# ---- file loaders ---------------------------------------------------------


def _read_csv_with_fallback(path: Path) -> pd.DataFrame:
    """Try common encodings before giving up. Excel-exported CSVs often
    aren't utf-8."""
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return pd.read_csv(path, encoding=encoding)
        except UnicodeDecodeError:
            continue
    # Last try: utf-8 with replacement
    return pd.read_csv(path, encoding="utf-8", encoding_errors="replace")


def load_file(path: Union[str, Path]) -> LoadResult:
    """Dispatch a file load by extension. Returns DataFrame for tabular,
    list[Document] (one element) for unstructured."""
    path = Path(path)
    ext = path.suffix.lower()

    if ext in {".csv", ".tsv"}:
        sep = "\t" if ext == ".tsv" else ","
        df = _read_csv_with_fallback(path) if ext == ".csv" else pd.read_csv(path, sep=sep)
        return df

    if ext in {".xlsx", ".xls"}:
        return pd.read_excel(path)

    if ext == ".pdf":
        text = _extract_pdf(path)
        return [Document(id=path.stem, text=text, source=str(path))]

    if ext == ".docx":
        text = _extract_docx(path)
        return [Document(id=path.stem, text=text, source=str(path))]

    if ext == ".txt":
        text = path.read_text(encoding="utf-8", errors="replace")
        return [Document(id=path.stem, text=text, source=str(path))]

    if ext in {".json", ".jsonl"}:
        # JSON returns a tabular DataFrame after path resolution. The
        # default path "." works for plain arrays of objects; nested
        # data needs the user to specify json_path through the UI.
        from . import json_loader
        return json_loader.load_json(path, json_path=".")

    raise ValueError(f"Unsupported file type {ext!r}. Supported: {SUPPORTED_FILE_EXTS}")


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError as exc:
        raise ImportError("pypdf required for PDF ingestion. pip install pypdf") from exc
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(path: Path) -> str:
    try:
        import docx  # python-docx  # type: ignore
    except ImportError as exc:
        raise ImportError(
            "python-docx required for DOCX. pip install python-docx"
        ) from exc
    doc = docx.Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def load_json_with_path(path: Union[str, Path], json_path: str = ".") -> pd.DataFrame:
    """Load a JSON or JSONL file, resolving `json_path` to find the records.

    Use this instead of load_file() when the user has supplied a custom
    path (via the UI's path-picker)."""
    from . import json_loader
    return json_loader.load_json(path, json_path=json_path)


def load_files(paths: list[Union[str, Path]]) -> dict[str, LoadResult]:
    """Load multiple files. Returns {path_str: result} so the caller can
    handle each file's structure (tabular vs documents) independently."""
    return {str(p): load_file(p) for p in paths}


def column_candidates(path: Union[str, Path]) -> list[str]:
    """For tabular files, return column names so the UI can offer pickers."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".csv":
        return list(_read_csv_with_fallback(path).columns)
    if ext == ".tsv":
        return list(pd.read_csv(path, sep="\t", nrows=0).columns)
    if ext in {".xlsx", ".xls"}:
        return list(pd.read_excel(path, nrows=0).columns)
    return []


# ---- paste ----------------------------------------------------------------


def load_paste(text: str, mode: str = "single") -> list[Document]:
    """Pasted text. mode='single' -> one doc; mode='lines' -> one per line."""
    if not text or not text.strip():
        return []
    if mode == "single":
        return [Document(id="paste_0", text=text, source="paste")]
    if mode == "lines":
        rows = []
        for i, line in enumerate(text.splitlines()):
            line = line.strip()
            if line:
                rows.append(Document(id=f"paste_{i}", text=line, source="paste_lines"))
        return rows
    raise ValueError(f"unknown paste mode {mode!r}")


# ---- HuggingFace ----------------------------------------------------------


def load_huggingface(
    name: str,
    *,
    split: str = "train",
    config: str | None = None,
    max_rows: int | None = None,
) -> pd.DataFrame:
    """Load a HF dataset by name. Returns a DataFrame.

    Examples:
        load_huggingface("imdb")
        load_huggingface("ag_news", split="test")
        load_huggingface("squad_v2", split="validation", max_rows=100)
    """
    try:
        from datasets import load_dataset  # type: ignore
    except ImportError as exc:
        raise ImportError(
            "datasets is required for HuggingFace ingestion. "
            "pip install datasets"
        ) from exc
    ds = load_dataset(name, config, split=split) if config else load_dataset(name, split=split)
    if max_rows is not None:
        ds = ds.select(range(min(max_rows, len(ds))))
    return ds.to_pandas()
